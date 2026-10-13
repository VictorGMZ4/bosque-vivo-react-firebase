"""Build the downloadable, blank Flora Viva Office templates. Requires python-docx/openpyxl."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

DEST = Path(__file__).resolve().parents[1] / 'public' / 'templates'
DEST.mkdir(parents=True, exist_ok=True)
FIELDS = ['Nombre común *', 'Nombre científico *', 'Familia botánica *', 'Género', 'Orden',
          'Otros nombres', 'Tipo de planta *', 'Descripción general', 'Características morfológicas',
          'Ecología y cultivo', 'Propiedades y usos', 'Precauciones', 'Fuentes y referencias']
GREEN = '244C39'
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.65)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(5)
doc.styles['Title'].font.color.rgb = RGBColor(0,0,0)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
sec.header.paragraphs[0].text = 'FLORA VIVA · BOSQUE VERDE  |  UNIPAZ'
doc.add_heading('Ficha para registrar una especie', 0)
doc.add_paragraph('Completa la columna Información con los datos documentados de una sola planta. Conserva los nombres de los campos y guarda el archivo como Word (.docx). Los campos con * son obligatorios.')
doc.add_paragraph('Tipo de planta: Árbol, Arbusto, Herbácea, Palmera o Trepadora. Deja vacíos los datos desconocidos. No añadas imágenes; podrás adjuntarlas después de guardar la especie en la web.')
for section, labels in [('Identificación y taxonomía', FIELDS[:7]), ('Descripción y fuentes', FIELDS[7:])]:
    if section == 'Descripción y fuentes': doc.add_page_break()
    doc.add_heading(section, 1)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    table.columns[0].width = Inches(2.1)
    table.columns[1].width = Inches(4.3)
    for cell, text in zip(table.rows[0].cells, ['Campo', 'Información']):
        cell.text = text
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), GREEN); cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs: run.font.color.rgb = RGBColor(255,255,255); run.bold = True
    for label in labels:
        row = table.add_row(); row.cells[0].text = label
        row.height = Inches(.43 if section == 'Identificación y taxonomía' else .87)
        props = row._tr.get_or_add_trPr(); props.append(OxmlElement('w:cantSplit'))
    for row in table.rows:
        row.cells[0].width = Inches(2.1); row.cells[1].width = Inches(4.3)
doc.add_paragraph('Al subir esta ficha, revisa los datos en Registrar especie y pulsa Guardar especie o Enviar a revisión. Subir el documento no publica la información.')
sec.footer.paragraphs[0].text = 'Flora Viva · Bosque Verde    |    Plantilla de importación v1'
doc.save(DEST / 'flora-viva-especie.docx')

book = Workbook(); sheet = book.active; sheet.title = 'Especies'
sheet.append(FIELDS); sheet.freeze_panes = 'D2'; sheet.auto_filter.ref = 'A1:M101'
for cell in sheet[1]:
    cell.fill = PatternFill('solid', fgColor=GREEN); cell.font = Font(color='FFFFFF', bold=True)
    cell.alignment = Alignment(wrap_text=True, vertical='center')
sheet.row_dimensions[1].height = 34
for col in range(1,14):
    from openpyxl.utils import get_column_letter
    sheet.column_dimensions[get_column_letter(col)].width = 26 if col < 8 else 45
    for row in range(2,102):
        cell = sheet.cell(row,col); cell.number_format = '@'
        cell.font = Font(color='174AA0', size=11)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        if row % 2 == 0: cell.fill = PatternFill('solid', fgColor='F0F6F1')
for row in range(2,102): sheet.row_dimensions[row].height = 42
validation = DataValidation(type='list', formula1='"Árbol,Arbusto,Herbácea,Palmera,Trepadora"', allow_blank=False)
validation.errorTitle = 'Selecciona un tipo de planta'; validation.error = 'Usa una opción de la lista.'; validation.showErrorMessage = True
sheet.add_data_validation(validation); validation.add('G2:G101')
sheet['A1'].comment = Comment('Una especie por fila. Completa los campos con *. No cambies los encabezados. Máximo 100 especies.', 'Flora Viva')
guide = book.create_sheet('Instrucciones')
instructions = [
 'Flora Viva · Bosque Verde — Plantilla de especies',
 '1. Rellena la hoja Especies: una planta por fila, hasta 100 plantas.',
 '2. Conserva la fila de encabezados y el nombre de la hoja Especies.',
 '3. Completa nombre común, nombre científico, familia y tipo de planta.',
 '4. Escribe texto; no uses fórmulas, imágenes ni vínculos a otros libros.',
 '5. Deja vacíos los datos desconocidos e incluye tus fuentes y referencias.',
 '6. Guarda como Excel (.xlsx) y sube el archivo en Registrar especie.',
 '7. Revisa cada fila en el formulario. Guardar registra solo la especie visible.',
 'Los profesores envían propuestas a revisión. La importación no cambia los permisos.',
 'Límites: nombres 140/180 caracteres; familia, género y orden 100; otros nombres 240.',
 'Cada descripción admite hasta 20 000 caracteres. Archivo máximo: 5 MB.',
 'Las fotos se añaden desde la ficha después de crear la especie.'
]
for line in instructions: guide.append([line])
guide.column_dimensions['A'].width = 110
for row in guide:
    row[0].alignment = Alignment(wrap_text=True, vertical='center'); guide.row_dimensions[row[0].row].height = 34
guide['A1'].font = Font(bold=True, color=GREEN, size=16)
book.save(DEST / 'flora-viva-especies.xlsx')
print(DEST)
